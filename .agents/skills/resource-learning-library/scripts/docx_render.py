#!/usr/bin/env python3
"""Render the supported Markdown subset to a themed DOCX and self-contained HTML."""
from __future__ import annotations

import argparse
import html
import re
import sys
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

NAVY = "0B3D5C"
BLUE = "12628F"
PALE = "F6FAFD"
INK = "183243"
MUTED = "536B7B"
BORDER = "C7D7E1"
WHITE = "FFFFFF"

_HEADING_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")
_LIST_RE = re.compile(r"^(\s*)([-+*]|\d+[.)])\s+(.*)$")
_FENCE_RE = re.compile(r"^\s*```\s*([^`]*)$")
_HR_RE = re.compile(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$")
_INLINE_RE = re.compile(
    r"(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|(?<!\*)\*[^*\n]+\*(?!\*)|`[^`\n]+`)")
_LINK_RE = re.compile(r"^\[([^\]]+)\]\(([^)]+)\)$")
_SHORTKEY_RE = re.compile(r"(?<![!\\])\[([A-Z][A-Z0-9_-]{1,19})\](?!\()")


def _strip_shortkeys(text: str) -> str:
    """Remove standalone uppercase cross-reference tags such as [HR1]."""
    return _SHORTKEY_RE.sub("", text)


def _split_table_row(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|") and not value.endswith(r"\|"):
        value = value[:-1]
    cells = re.split(r"(?<!\\)\|", value)
    return [cell.replace(r"\|", "|").strip() for cell in cells]


def _is_table_delimiter(line: str) -> bool:
    if "|" not in line:
        return False
    cells = _split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def _starts_block(lines: list[str], index: int) -> bool:
    line = lines[index]
    stripped = line.lstrip()
    if not stripped:
        return True
    if _HEADING_RE.match(line) or _FENCE_RE.match(line) or _LIST_RE.match(line):
        return True
    if stripped.startswith(">") or _HR_RE.match(line):
        return True
    if index + 1 < len(lines) and _is_table_delimiter(lines[index + 1]) and "|" in line:
        return True
    return False


def parse(md: str) -> list[dict[str, Any]]:
    """Parse a practical Markdown subset into renderer-neutral block dictionaries."""
    lines = md.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[dict[str, Any]] = []
    i = 0

    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue

        fence = _FENCE_RE.match(line)
        if fence:
            language = fence.group(1).strip()
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not re.match(r"^\s*```\s*$", lines[i]):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            blocks.append({"type": "code", "language": language, "text": "\n".join(code_lines)})
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            blocks.append({"type": "heading", "level": len(heading.group(1)), "text": heading.group(2).strip()})
            i += 1
            continue

        if i + 1 < len(lines) and _is_table_delimiter(lines[i + 1]) and "|" in line:
            table_rows = [_split_table_row(line)]
            i += 2  # header row and separator
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                table_rows.append(_split_table_row(lines[i]))
                i += 1
            width = max((len(row) for row in table_rows), default=0)
            normalized = [row + [""] * (width - len(row)) for row in table_rows]
            blocks.append({"type": "table", "rows": normalized})
            continue

        if _HR_RE.match(line):
            blocks.append({"type": "hr"})
            i += 1
            continue

        if line.lstrip().startswith(">"):
            quote_lines: list[str] = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote_lines.append(re.sub(r"^\s{0,3}>\s?", "", lines[i]))
                i += 1
            text = "\n".join(quote_lines).strip()
            variant = "note"
            marker = re.match(r"^\[!(NOTE|TIP|WARNING|IMPORTANT)\]\s*", text, re.I)
            if marker:
                variant = marker.group(1).lower()
                text = text[marker.end():]
            else:
                label = re.match(r"^(Tip|Note|Warning|Important):\s*", text, re.I)
                if label:
                    variant = label.group(1).lower()
            blocks.append({"type": "callout", "variant": variant, "text": text})
            continue

        list_match = _LIST_RE.match(line)
        if list_match:
            items: list[dict[str, Any]] = []
            while i < len(lines):
                item_match = _LIST_RE.match(lines[i])
                if not item_match:
                    break
                indent = len(item_match.group(1).expandtabs(2))
                marker = item_match.group(2)
                items.append({
                    "text": item_match.group(3).strip(),
                    "ordered": marker[0].isdigit(),
                    "level": max(0, indent // 2),
                })
                i += 1
            blocks.append({"type": "list", "items": items})
            continue

        paragraph_lines = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not _starts_block(lines, i):
            paragraph_lines.append(lines[i].strip())
            i += 1
        blocks.append({"type": "paragraph", "text": " ".join(paragraph_lines)})

    return blocks


def _safe_url(raw_url: str) -> str | None:
    url = html.unescape(raw_url.strip())
    parts = urlsplit(url)
    if parts.scheme in {"http", "https"} and parts.netloc:
        return url
    if parts.scheme == "mailto" and parts.path:
        return url
    return None


def _first_h1(blocks: list[dict[str, Any]]) -> str | None:
    for block in blocks:
        if block.get("type") == "heading" and block.get("level") == 1:
            return _strip_shortkeys(str(block.get("text", ""))).strip()
    return None


def _without_first_h1(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    removed = False
    for block in blocks:
        if not removed and block.get("type") == "heading" and block.get("level") == 1:
            removed = True
            continue
        result.append(block)
    return result


def _hero_meta_text(hero_meta: Any) -> str:
    if hero_meta is None:
        return ""
    if isinstance(hero_meta, str):
        return hero_meta.strip()
    if isinstance(hero_meta, (list, tuple)):
        return "  ·  ".join(str(item).strip() for item in hero_meta if str(item).strip())
    return str(hero_meta).strip()


def _inline_html(text: str) -> str:
    text = _strip_shortkeys(text)
    chunks: list[str] = []
    cursor = 0
    for match in _INLINE_RE.finditer(text):
        chunks.append(html.escape(text[cursor:match.start()]))
        token = match.group(0)
        link = _LINK_RE.match(token)
        if link:
            label, raw_url = link.groups()
            safe_url = _safe_url(raw_url)
            if safe_url:
                chunks.append(f'<a href="{html.escape(safe_url, quote=True)}">{html.escape(_strip_shortkeys(label))}</a>')
            else:
                chunks.append(html.escape(label))
        elif token.startswith("**") and token.endswith("**"):
            chunks.append(f"<strong>{html.escape(_strip_shortkeys(token[2:-2]))}</strong>")
        elif token.startswith("*") and token.endswith("*"):
            chunks.append(f"<em>{html.escape(_strip_shortkeys(token[1:-1]))}</em>")
        elif token.startswith("`") and token.endswith("`"):
            chunks.append(f"<code>{html.escape(_strip_shortkeys(token[1:-1]))}</code>")
        else:
            chunks.append(html.escape(token))
        cursor = match.end()
    chunks.append(html.escape(text[cursor:]))
    return "".join(chunks)


def _add_hyperlink(paragraph: Any, label: str, url: str, color: str = BLUE) -> None:
    """Add a real external hyperlink relationship to a python-docx paragraph."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color or BLUE)
    rpr.append(color_element)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    rpr.append(fonts)
    run.append(rpr)
    text_element = OxmlElement("w:t")
    text_element.text = _strip_shortkeys(label)
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_inline_docx(paragraph: Any, text: str, *, base_size: float | None = None, base_color: str | None = None) -> None:
    from docx.oxml import OxmlElement

    text = _strip_shortkeys(text)
    cursor = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > cursor:
            _add_plain_run(paragraph, text[cursor:match.start()], base_size=base_size, base_color=base_color)
        token = match.group(0)
        link = _LINK_RE.match(token)
        if link:
            label, raw_url = link.groups()
            safe_url = _safe_url(raw_url)
            if safe_url:
                link_color = WHITE if (base_color or "").upper() == WHITE else BLUE
                _add_hyperlink(paragraph, label, safe_url, color=link_color)
            else:
                _add_plain_run(paragraph, label, base_size=base_size, base_color=base_color)
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(_strip_shortkeys(token[2:-2]))
            run.bold = True
            _style_run(run, size=base_size, color=base_color)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(_strip_shortkeys(token[1:-1]))
            run.italic = True
            _style_run(run, size=base_size, color=base_color)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(_strip_shortkeys(token[1:-1]))
            run.font.name = "Consolas"
            _style_run(run, size=base_size or 9, color=base_color or BLUE)
        else:
            _add_plain_run(paragraph, token, base_size=base_size, base_color=base_color)
        cursor = match.end()
    if cursor < len(text):
        _add_plain_run(paragraph, text[cursor:], base_size=base_size, base_color=base_color)


def _style_run(run: Any, *, size: float | None = None, color: str | None = None) -> None:
    from docx.shared import Pt, RGBColor

    run.font.name = "Calibri"
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_plain_run(paragraph: Any, text: str, *, base_size: float | None = None, base_color: str | None = None) -> None:
    parts = text.split("\n")
    for index, part in enumerate(parts):
        if part:
            run = paragraph.add_run(part)
            _style_run(run, size=base_size, color=base_color)
        if index < len(parts) - 1:
            paragraph.add_run().add_break()


def _set_cell_shading(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:shd"))
    if old is not None:
        tc_pr.remove(old)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_cell_margins(cell: Any, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table: Any, color: str = BORDER, size: int = 5) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _shade_paragraph(paragraph: Any, fill: str, border: str | None = None) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    if border:
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        for side in ("left", "top", "bottom", "right"):
            edge = borders.find(qn(f"w:{side}"))
            if edge is None:
                edge = OxmlElement(f"w:{side}")
                borders.append(edge)
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "18" if side == "left" else "4")
            edge.set(qn("w:space"), "5")
            edge.set(qn("w:color"), border)
        ind = p_pr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            p_pr.append(ind)
        ind.set(qn("w:left"), "120")
        ind.set(qn("w:right"), "120")


def _add_callout(doc: Any, block: dict[str, Any]) -> None:
    from docx.shared import Pt, RGBColor

    variant = str(block.get("variant", "note")).lower()
    palette = {
        "note": (PALE, BLUE, INK),
        "tip": ("EAF5F7", "168078", INK),
        "warning": ("FFF5E8", "B56A12", INK),
        "important": ("FCEEEF", "A63B4B", INK),
    }
    fill, border, ink = palette.get(variant, palette["note"])
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_together = True
    _shade_paragraph(paragraph, fill, border)
    text = str(block.get("text", ""))
    label_match = re.match(r"^(Tip|Note|Warning|Important):\s*", text, re.I)
    if label_match:
        label = paragraph.add_run(label_match.group(1).title() + ": ")
        label.bold = True
        label.font.color.rgb = RGBColor.from_string(border)
        text = text[label_match.end():]
    _add_inline_docx(paragraph, text, base_size=10, base_color=ink)


def _add_code(doc: Any, block: dict[str, Any]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    lines = str(block.get("text", "")).split("\n") or [""]
    for index, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        p_pr = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), PALE)
        p_pr.append(shading)
        p_bdr = OxmlElement("w:pBdr")
        for side in ("left", "right", "top", "bottom"):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "5")
            edge.set(qn("w:space"), "3")
            edge.set(qn("w:color"), BORDER)
            if side in {"top", "bottom"} and ((side == "top" and index > 0) or (side == "bottom" and index < len(lines) - 1)):
                edge.set(qn("w:val"), "nil")
            p_bdr.append(edge)
        p_pr.append(p_bdr)
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(INK)


def _add_table(doc: Any, rows: list[list[str]]) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.shared import Pt, RGBColor

    if not rows:
        return
    width = max((len(row) for row in rows), default=1)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    _set_table_borders(table)

    for row_index, row in enumerate(rows):
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if row_index == 0:
                _set_cell_shading(cell, NAVY)
            elif row_index % 2 == 1:
                _set_cell_shading(cell, PALE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            value = row[col_index] if col_index < len(row) else ""
            _add_inline_docx(paragraph, str(value), base_size=8.7, base_color=WHITE if row_index == 0 else INK)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _add_footer(section: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    _shade_paragraph(paragraph, PALE)
    run = paragraph.add_run("RESOURCE LEARNING LIBRARY  ·  ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    field_rpr = OxmlElement("w:rPr")
    field_color = OxmlElement("w:color")
    field_color.set(qn("w:val"), MUTED)
    field_rpr.append(field_color)
    field_size = OxmlElement("w:sz")
    field_size.set(qn("w:val"), "16")
    field_rpr.append(field_size)
    field_run.append(field_rpr)
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    paragraph._p.append(field)


def build_docx(
    blocks: list[dict[str, Any]],
    path: str | Path,
    title: str | None = None,
    subtitle: str | None = None,
    skip_first_h1: bool = False,
    hero_meta: Any = None,
) -> Path:
    """Build a styled Word document from parsed Markdown blocks."""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("python-docx is required; install it with `python3 -m pip install python-docx`.") from exc

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in (("Heading 1", 20, NAVY), ("Heading 2", 15, BLUE), ("Heading 3", 12, BLUE)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    heading_title = title or _first_h1(blocks)
    if heading_title:
        hero = doc.add_paragraph()
        hero.paragraph_format.space_before = Pt(3)
        hero.paragraph_format.space_after = Pt(4)
        hero.paragraph_format.keep_with_next = True
        _shade_paragraph(hero, NAVY)
        p_pr = hero._p.get_or_add_pPr()
        indent = p_pr.find(qn("w:ind"))
        if indent is None:
            indent = OxmlElement("w:ind")
            p_pr.append(indent)
        indent.set(qn("w:left"), "180")
        indent.set(qn("w:right"), "180")
        spacing = p_pr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            p_pr.append(spacing)
        spacing.set(qn("w:before"), "120")
        spacing.set(qn("w:after"), "120")
        title_run = hero.add_run(_strip_shortkeys(heading_title))
        title_run.font.name = "Calibri"
        title_run.font.size = Pt(23)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor.from_string(WHITE)

    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(_strip_shortkeys(subtitle))
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string(BLUE)
    meta = _hero_meta_text(hero_meta)
    if meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(9)
        run = p.add_run(_strip_shortkeys(meta))
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    source_blocks = _without_first_h1(blocks) if skip_first_h1 else blocks
    for block in source_blocks:
        kind = block.get("type")
        if kind == "heading":
            level = max(1, min(3, int(block.get("level", 1))))
            p = doc.add_paragraph(style=f"Heading {level}")
            _add_inline_docx(p, str(block.get("text", "")))
        elif kind == "paragraph":
            p = doc.add_paragraph()
            _add_inline_docx(p, str(block.get("text", "")))
        elif kind == "list":
            for item in block.get("items", []):
                style_name = "List Number" if item.get("ordered") else "List Bullet"
                p = doc.add_paragraph(style=style_name)
                p.paragraph_format.space_after = Pt(3)
                level = min(int(item.get("level", 0)), 4)
                if level:
                    p.paragraph_format.left_indent = Inches(0.25 * level)
                _add_inline_docx(p, str(item.get("text", "")))
        elif kind == "table":
            _add_table(doc, block.get("rows", []))
        elif kind == "callout":
            _add_callout(doc, block)
        elif kind == "code":
            _add_code(doc, block)
        elif kind == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            p_pr = p._p.get_or_add_pPr()
            p_bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), BORDER)
            p_bdr.append(bottom)
            p_pr.append(p_bdr)

    _add_footer(section)
    doc.core_properties.title = heading_title or output.stem
    doc.core_properties.subject = "Resource-learning library"
    doc.core_properties.author = "Resource Learning Library skill"
    doc.save(output)
    return output


_HTML_CSS = """
:root { color-scheme: light; --navy:#0B3D5C; --blue:#12628F; --pale:#F6FAFD; --ink:#183243; --muted:#536B7B; --border:#C7D7E1; }
* { box-sizing: border-box; }
body { margin: 0; padding: 28px 18px 48px; background: #f4f7f9; color: var(--ink); font: 16px/1.55 system-ui, -apple-system, "Segoe UI", sans-serif; }
main { max-width: 980px; margin: 0 auto; padding: 32px clamp(20px, 5vw, 56px); background: #fff; border: 1px solid #e1e9ee; border-radius: 16px; box-shadow: 0 12px 36px rgba(11,61,92,.08); }
.hero { margin: -32px calc(-1 * clamp(20px, 5vw, 56px)) 26px; padding: 26px clamp(20px, 5vw, 56px); color: #fff; background: var(--navy); border-radius: 16px 16px 0 0; }
.hero h1 { margin: 0; color: #fff; font-size: clamp(1.8rem, 4vw, 2.5rem); line-height: 1.15; }
.hero .subtitle { margin: 10px 0 0; color: #e4f0f6; font-size: 1.1rem; }
.hero .meta { margin: 8px 0 0; color: #c9dce7; font-size: .9rem; }
h1, h2, h3, h4, h5, h6 { color: var(--navy); line-height: 1.25; margin: 1.55em 0 .55em; }
h2, h4, h6 { color: var(--blue); }
p { margin: .55em 0 .9em; }
a { color: var(--blue); text-decoration-thickness: 1px; text-underline-offset: 2px; overflow-wrap: anywhere; }
table { width: 100%; border-collapse: collapse; margin: 1.1em 0 1.5em; font-size: .93rem; }
th, td { border: 1px solid var(--border); padding: 9px 11px; text-align: left; vertical-align: top; }
th { background: var(--navy); color: #fff; }
tbody tr:nth-child(even) { background: var(--pale); }
.callout { margin: 1em 0; padding: 12px 16px; border-left: 5px solid var(--blue); background: var(--pale); border-radius: 4px; }
.callout.tip { border-color: #168078; background: #eaf5f7; }
.callout.warning { border-color: #b56a12; background: #fff5e8; }
.callout.important { border-color: #a63b4b; background: #fceeef; }
pre { overflow-x: auto; padding: 14px 16px; border: 1px solid var(--border); border-radius: 8px; background: var(--pale); font: .88rem/1.5 Consolas, "SFMono-Regular", monospace; }
code { padding: .1em .3em; background: var(--pale); border-radius: 4px; font-family: Consolas, "SFMono-Regular", monospace; }
pre code { padding: 0; background: transparent; }
li { margin: .28em 0; }
hr { border: 0; border-top: 1px solid var(--border); margin: 1.5em 0; }
@media print { body { padding: 0; background: white; } main { max-width: none; border: 0; box-shadow: none; } .hero { border-radius: 0; } a { color: var(--blue); } }
"""


def build_html(
    blocks: list[dict[str, Any]],
    path: str | Path,
    title: str | None = None,
    subtitle: str | None = None,
    skip_first_h1: bool = False,
    hero_meta: Any = None,
) -> Path:
    """Build a self-contained HTML document using inline CSS and no external assets."""
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    heading_title = title or _first_h1(blocks)
    source_blocks = _without_first_h1(blocks) if skip_first_h1 else blocks
    body: list[str] = []
    if heading_title:
        hero = [f"<h1>{html.escape(_strip_shortkeys(heading_title))}</h1>"]
        if subtitle:
            hero.append(f'<p class="subtitle">{html.escape(_strip_shortkeys(subtitle))}</p>')
        meta = _hero_meta_text(hero_meta)
        if meta:
            hero.append(f'<p class="meta">{html.escape(_strip_shortkeys(meta))}</p>')
        body.append('<header class="hero">' + "\n".join(hero) + "</header>")

    for block in source_blocks:
        kind = block.get("type")
        if kind == "heading":
            level = max(1, min(6, int(block.get("level", 1))))
            body.append(f"<h{level}>{_inline_html(str(block.get('text', '')))}</h{level}>")
        elif kind == "paragraph":
            body.append(f"<p>{_inline_html(str(block.get('text', '')))}</p>")
        elif kind == "list":
            items = block.get("items", [])
            if items:
                groups: list[tuple[bool, list[dict[str, Any]]]] = []
                for item in items:
                    ordered = bool(item.get("ordered"))
                    if not groups or groups[-1][0] != ordered:
                        groups.append((ordered, []))
                    groups[-1][1].append(item)
                for ordered, group_items in groups:
                    tag = "ol" if ordered else "ul"
                    li_html = []
                    for item in group_items:
                        depth = max(0, min(8, int(item.get("level", 0))))
                        style = f' style="margin-left:{depth * 1.25}em"' if depth else ""
                        li_html.append(f"<li{style}>{_inline_html(str(item.get('text', '')))}</li>")
                    body.append(f"<{tag}>" + "".join(li_html) + f"</{tag}>")
        elif kind == "table":
            rows = block.get("rows", [])
            if rows:
                header = rows[0]
                body_rows = rows[1:]
                ths = "".join(f"<th scope=\"col\">{_inline_html(str(cell))}</th>" for cell in header)
                trs = []
                for row in body_rows:
                    cells = "".join(f"<td>{_inline_html(str(cell))}</td>" for cell in row)
                    trs.append(f"<tr>{cells}</tr>")
                body.append(f"<table><thead><tr>{ths}</tr></thead><tbody>{''.join(trs)}</tbody></table>")
        elif kind == "callout":
            variant = str(block.get("variant", "note")).lower()
            safe_variant = variant if variant in {"note", "tip", "warning", "important"} else "note"
            body.append(f'<aside class="callout {safe_variant}">{_inline_html(str(block.get("text", "")))}</aside>')
        elif kind == "code":
            language = str(block.get("language", "")).strip()
            code_class = f' class="language-{html.escape(language, quote=True)}"' if language else ""
            body.append(f"<pre><code{code_class}>{html.escape(str(block.get('text', '')))}</code></pre>")
        elif kind == "hr":
            body.append("<hr>")

    title_text = heading_title or "Resource-learning library"
    document = (
        "<!doctype html>\n<html lang=\"en\">\n<head>\n"
        "<meta charset=\"utf-8\">\n<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
        f"<title>{html.escape(_strip_shortkeys(title_text))}</title>\n<style>\n{_HTML_CSS}\n</style>\n"
        "</head>\n<body>\n<main>\n" + "\n".join(body) + "\n</main>\n</body>\n</html>\n"
    )
    output.write_text(document, encoding="utf-8")
    return output


def _safe_output_name(name: str) -> str:
    """Reduce a user-supplied output name to a safe single filename stem."""
    candidate = Path(name).name
    if candidate.lower().endswith((".docx", ".html", ".htm")):
        candidate = candidate.rsplit(".", 1)[0]
    candidate = re.sub(r"[^\w.-]+", "-", candidate, flags=re.UNICODE).strip(" .-")
    return candidate or "resource-library"


def _reportlab_inline(text: str) -> str:
    """Convert supported inline Markdown to ReportLab Paragraph markup."""
    from xml.sax.saxutils import escape as xml_escape

    text = _strip_shortkeys(text)
    chunks: list[str] = []
    cursor = 0
    for match in _INLINE_RE.finditer(text):
        chunks.append(xml_escape(text[cursor:match.start()].replace("\n", "<br/>")))
        token = match.group(0)
        link = _LINK_RE.match(token)
        if link:
            label, raw_url = link.groups()
            safe_url = _safe_url(raw_url)
            if safe_url:
                href = html.escape(safe_url, quote=True)
                chunks.append(f'<link href="{href}">{xml_escape(_strip_shortkeys(label))}</link>')
            else:
                chunks.append(xml_escape(label))
        elif token.startswith("**") and token.endswith("**"):
            chunks.append(f"<b>{xml_escape(_strip_shortkeys(token[2:-2]))}</b>")
        elif token.startswith("*") and token.endswith("*"):
            chunks.append(f"<i>{xml_escape(_strip_shortkeys(token[1:-1]))}</i>")
        elif token.startswith("`") and token.endswith("`"):
            chunks.append(f'<font name="Courier" backColor="#{PALE}">{xml_escape(_strip_shortkeys(token[1:-1]))}</font>')
        else:
            chunks.append(xml_escape(token))
        cursor = match.end()
    chunks.append(xml_escape(text[cursor:].replace("\n", "<br/>")))
    return "".join(chunks)


def build_pdf(
    blocks: list[dict[str, Any]],
    path: str | Path,
    title: str | None = None,
    subtitle: str | None = None,
    skip_first_h1: bool = False,
    hero_meta: Any = None,
) -> Path:
    """Build a formatted PDF from parsed Markdown blocks using ReportLab."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import HRFlowable, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError("PDF output requires ReportLab; install it with `python3 -m pip install reportlab`.") from exc

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    heading_title = title or _first_h1(blocks) or output.stem
    source_blocks = _without_first_h1(blocks) if skip_first_h1 else blocks
    navy = colors.HexColor(f"#{NAVY}")
    blue = colors.HexColor(f"#{BLUE}")
    pale = colors.HexColor(f"#{PALE}")
    ink = colors.HexColor(f"#{INK}")
    muted = colors.HexColor(f"#{MUTED}")
    border = colors.HexColor(f"#{BORDER}")

    doc = SimpleDocTemplate(
        str(output), pagesize=A4,
        rightMargin=0.62 * inch, leftMargin=0.62 * inch,
        topMargin=0.62 * inch, bottomMargin=0.68 * inch,
        title=heading_title, author="Resource Learning Library skill",
    )
    usable_width = A4[0] - 1.24 * inch
    sample = getSampleStyleSheet()["BodyText"]
    body_style = ParagraphStyle("RLLBody", parent=sample, fontName="Helvetica", fontSize=9.5, leading=13.2, textColor=ink, spaceAfter=7)
    title_style = ParagraphStyle("RLLTitle", parent=sample, fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.white, alignment=TA_LEFT)
    subtitle_style = ParagraphStyle("RLLSubtitle", parent=sample, fontName="Helvetica", fontSize=12, leading=16, textColor=blue, spaceAfter=3)
    meta_style = ParagraphStyle("RLLMeta", parent=sample, fontName="Helvetica", fontSize=8.5, leading=11, textColor=muted, spaceAfter=10)
    header_cell_style = ParagraphStyle("RLLTableHeader", parent=body_style, fontName="Helvetica-Bold", textColor=colors.white, fontSize=8.2, leading=10.2, spaceAfter=0)
    table_cell_style = ParagraphStyle("RLLTableCell", parent=body_style, fontSize=8.2, leading=10.5, spaceAfter=0)
    code_style = ParagraphStyle("RLLCode", parent=sample, fontName="Courier", fontSize=7.5, leading=9.5, textColor=ink, leftIndent=5, rightIndent=5, spaceBefore=4, spaceAfter=4)
    heading_styles = {
        1: ParagraphStyle("RLLH1", parent=sample, fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=navy, spaceBefore=15, spaceAfter=7, keepWithNext=True),
        2: ParagraphStyle("RLLH2", parent=sample, fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=blue, spaceBefore=13, spaceAfter=6, keepWithNext=True),
        3: ParagraphStyle("RLLH3", parent=sample, fontName="Helvetica-Bold", fontSize=11.5, leading=15, textColor=blue, spaceBefore=10, spaceAfter=5, keepWithNext=True),
    }
    story: list[Any] = []

    title_table = Table([[Paragraph(html.escape(_strip_shortkeys(heading_title)), title_style)]], colWidths=[usable_width])
    title_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), navy),
        ("LEFTPADDING", (0, 0), (-1, -1), 14), ("RIGHTPADDING", (0, 0), (-1, -1), 14),
        ("TOPPADDING", (0, 0), (-1, -1), 14), ("BOTTOMPADDING", (0, 0), (-1, -1), 14),
    ]))
    story.extend([title_table, Spacer(1, 8)])
    if subtitle:
        story.append(Paragraph(html.escape(_strip_shortkeys(subtitle)), subtitle_style))
    meta = _hero_meta_text(hero_meta)
    if meta:
        story.append(Paragraph(html.escape(_strip_shortkeys(meta)), meta_style))

    callout_palette = {
        "note": (pale, blue),
        "tip": (colors.HexColor("#EAF5F7"), colors.HexColor("#168078")),
        "warning": (colors.HexColor("#FFF5E8"), colors.HexColor("#B56A12")),
        "important": (colors.HexColor("#FCEEEF"), colors.HexColor("#A63B4B")),
    }
    for block in source_blocks:
        kind = block.get("type")
        if kind == "heading":
            level = max(1, min(3, int(block.get("level", 1))))
            story.append(Paragraph(_reportlab_inline(str(block.get("text", ""))), heading_styles[level]))
        elif kind == "paragraph":
            story.append(Paragraph(_reportlab_inline(str(block.get("text", ""))), body_style))
        elif kind == "list":
            ordered_index = 0
            previous_ordered: bool | None = None
            for item in block.get("items", []):
                ordered = bool(item.get("ordered"))
                ordered_index = ordered_index + 1 if ordered and previous_ordered else (1 if ordered else 0)
                previous_ordered = ordered
                depth = max(0, min(6, int(item.get("level", 0))))
                list_style = ParagraphStyle(
                    f"RLLList{depth}", parent=body_style,
                    leftIndent=18 + (depth * 14), firstLineIndent=0,
                    bulletIndent=4 + (depth * 14), spaceAfter=3,
                )
                bullet = f"{ordered_index}." if ordered else "•"
                story.append(Paragraph(_reportlab_inline(str(item.get("text", ""))), list_style, bulletText=bullet))
        elif kind == "table":
            rows = block.get("rows", [])
            if rows:
                width = max((len(row) for row in rows), default=1)
                data = []
                for row_index, row in enumerate(rows):
                    style = header_cell_style if row_index == 0 else table_cell_style
                    data.append([Paragraph(_reportlab_inline(str(cell)), style) for cell in row] + [""] * (width - len(row)))
                col_width = usable_width / width
                table = Table(data, colWidths=[col_width] * width, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), navy),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, pale]),
                    ("GRID", (0, 0), (-1, -1), 0.45, border),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]))
                story.extend([table, Spacer(1, 8)])
        elif kind == "callout":
            variant = str(block.get("variant", "note")).lower()
            fill, accent = callout_palette.get(variant, callout_palette["note"])
            callout = Table([[Paragraph(_reportlab_inline(str(block.get("text", ""))), body_style)]], colWidths=[usable_width])
            callout.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), fill),
                ("BOX", (0, 0), (-1, -1), 0.5, border),
                ("LINEBEFORE", (0, 0), (0, -1), 3, accent),
                ("LEFTPADDING", (0, 0), (-1, -1), 11), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]))
            story.extend([callout, Spacer(1, 8)])
        elif kind == "code":
            code = str(block.get("text", "")) or " "
            code_panel = Table([[Preformatted(code, code_style)]], colWidths=[usable_width])
            code_panel.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), pale),
                ("BOX", (0, 0), (-1, -1), 0.5, border),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([code_panel, Spacer(1, 8)])
        elif kind == "hr":
            story.extend([Spacer(1, 4), HRFlowable(width="100%", thickness=0.6, color=border), Spacer(1, 8)])

    def draw_footer(canvas: Any, _doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted)
        canvas.drawCentredString(A4[0] / 2, 0.38 * inch, f"Resource Learning Library  ·  Page {_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return output


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", help="Markdown source file")
    parser.add_argument("-o", "--output-dir", default=".", help="output directory (default: current directory)")
    parser.add_argument("--name", help="output base name (default: source filename)")
    parser.add_argument("--title", help="document title (defaults to the first H1, then output name)")
    parser.add_argument("--subtitle", help="optional subtitle for the title band")
    parser.add_argument("--hero-meta", help="optional short metadata line for the title band")
    parser.add_argument("--keep-first-h1", action="store_true", help="keep the first H1 in the body even when it is used as the title")
    parser.add_argument("--skip-first-h1", action="store_true", help="skip the first H1 in the body")
    outputs = parser.add_mutually_exclusive_group()
    outputs.add_argument("--format", nargs="+", choices=("pdf", "docx", "md", "html"), dest="formats", help="one or more output formats; default: docx html")
    outputs.add_argument("--docx-only", action="store_true", help="write only the DOCX (legacy convenience option)")
    outputs.add_argument("--pdf-only", action="store_true", help="write only the PDF")
    outputs.add_argument("--html-only", action="store_true", help="write only the static HTML")
    outputs.add_argument("--md-only", action="store_true", help="write only a copy of the Markdown source")
    args = parser.parse_args(argv)

    if args.keep_first_h1 and args.skip_first_h1:
        parser.error("--keep-first-h1 and --skip-first-h1 cannot be used together")

    source = Path(args.markdown)
    if not source.is_file():
        parser.error(f"Markdown source not found: {source}")
    try:
        markdown = source.read_text(encoding="utf-8")
    except OSError as exc:
        parser.error(f"Could not read {source}: {exc}")
    blocks = parse(markdown)
    base_name = _safe_output_name(args.name or source.stem)
    first_h1 = _first_h1(blocks)
    title = args.title or first_h1 or base_name
    skip_first = args.skip_first_h1 or (bool(first_h1) and not args.keep_first_h1)
    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.formats:
        formats = list(dict.fromkeys(args.formats))
    elif args.pdf_only:
        formats = ["pdf"]
    elif args.docx_only:
        formats = ["docx"]
    elif args.html_only:
        formats = ["html"]
    elif args.md_only:
        formats = ["md"]
    else:
        formats = ["docx", "html"]

    try:
        for output_format in formats:
            output_path = out_dir / f"{base_name}.{output_format}"
            if output_format == "docx":
                build_docx(blocks, output_path, title=title, subtitle=args.subtitle, skip_first_h1=skip_first, hero_meta=args.hero_meta)
            elif output_format == "pdf":
                build_pdf(blocks, output_path, title=title, subtitle=args.subtitle, skip_first_h1=skip_first, hero_meta=args.hero_meta)
            elif output_format == "html":
                build_html(blocks, output_path, title=title, subtitle=args.subtitle, skip_first_h1=skip_first, hero_meta=args.hero_meta)
            elif output_format == "md":
                if source.resolve() != output_path.resolve():
                    output_path.write_text(markdown, encoding="utf-8")
            print(f"Wrote {output_path}")
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
